#!/usr/bin/env python3
"""
One shareable workbook of everything the bot has done, live and paper.

    python export_trade_log.py             # fetch wallet history (cached) + build
    python export_trade_log.py --refresh   # re-fetch the wallet history first

Output: exports/robinhood-bot-trade-log.xlsx  (gitignored -- personal records)

Where each sheet comes from, and why:
  Live Transactions   the chain itself, via the Blockscout explorer. Ground truth
                      for every buy, sell, approval and failed transaction, and
                      the gas each one burned -- including failures the bot never
                      logged.
  Live Trades         the bot's log + trade ledger, LIVE sessions only.
  Paper Trades        the paper ledger, plus the DRY_RUN sessions of 09-12 that
                      wrote into the live log folder before paper mode had its own.
  Signals & Failures  every call received, rejected, or that failed to execute.

Deliberately excluded: the rows this project's own test suite wrote into the
live ledger (ticker "TEST", tx hash "0xfirst", repeated-byte fixture contracts
like 0x1111...). Found by signature, counted, and reported in Notes.

Wallet: WALLET_ADDRESS in .env, else inferred from the cached history.

Runs on system Python (needs openpyxl; the bot's venv is left untouched).
Read-only against the chain. Sends nothing.
"""
import argparse
import json
import os
import re
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
import zipfile
from collections import Counter
from datetime import datetime

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import dashboard_data as dd  # noqa: E402  (stdlib-only collector, reused)

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

BS = "https://robinhoodchain.blockscout.com/api/v2"
EXPLORER = "https://robinhoodchain.blockscout.com/tx/"
# A bare User-Agent gets 403 from this explorer; a full browser one does not.
HDR = {"User-Agent": ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                      "(KHTML, like Gecko) Chrome/126.0.0.0 Safari/537.36"),
       "Accept": "application/json, text/plain, */*",
       "Referer": "https://robinhoodchain.blockscout.com/"}

WALLET_CACHE = os.path.join(HERE, "cache", "wallet_txs.json")
OUT = os.path.join(HERE, "exports", "robinhood-bot-trade-log.xlsx")
LIVE_LEDGER = os.path.join(HERE, "logs", "events.jsonl")
PAPER_LEDGER = os.path.join(HERE, "logs", "paper", "events.jsonl")
PAPER_LOG = os.path.join(HERE, "logs", "paper", "bot.log")


def wallet_address():
    """The bot's wallet. Deliberately not written in this file: it is published,
    and a hardcoded address would tie the wallet to whoever publishes it.
    WALLET_ADDRESS (environment or .env) wins; otherwise the one address that is
    party to every transaction in the cached history -- which is the wallet."""
    v = os.environ.get("WALLET_ADDRESS", "")
    env = os.path.join(HERE, ".env")
    if not v and os.path.exists(env):
        for line in open(env, encoding="utf-8", errors="replace"):
            if line.startswith("WALLET_ADDRESS="):
                v = line.split("=", 1)[1].strip().strip("\"'")
    if not v and os.path.exists(WALLET_CACHE):
        parties = [{((t.get(k) or {}).get("hash") or "").lower() for k in ("from", "to")}
                   for t in json.load(open(WALLET_CACHE, encoding="utf-8"))]
        common = (set.intersection(*parties) - {""}) if parties else set()
        v = common.pop() if len(common) == 1 else ""
    if not re.fullmatch(r"0x[0-9a-fA-F]{40}", v):
        sys.exit("Set WALLET_ADDRESS=0x... in .env -- this file does not keep the address.")
    return v


WALLET = wallet_address()

QUOTES = {"0x" + "0" * 40,
          "0x0bd7d308f8e1639fab988df18a8011f41eacad73",    # WETH
          "0x5fc5360d0400a0fd4f2af552add042d716f1d168",    # USDG
          "0x92fd66527192e3e61d4ddd13322aa222de86f9b5"}    # SGOV

# 1inch and 0x routers. Nothing in this repo calls either (grepped), so a
# transaction through them came from a wallet app, not from the bot.
NOT_BOT_ROUTES = {"0x5a705de8982235a7fa45bb83dcacf03a211389c7": "1inch",
                  "0x0000000000001ff3684f28c67538d4d072c22734": "0x"}
CALL_ADDR = re.compile(r"for \$?([A-Za-z0-9_]{2,20})\. Addr: (0x[0-9a-fA-F]{40})")

# Known artifacts: rows that are genuinely in the record but whose numbers are
# not real, with where the reason is written down.
FLAGS = {"ONCHAIN_842B9C": "Price-feed bug: held blind 110 min, P&L not real (AD-034)"}

FONT = "Arial"
MONEY = '$#,##0.00;($#,##0.00);"-"'
ETH_FMT = '0.000000;(0.000000);"-"'
PCT = '0.0%;(0.0%);"-"'


# --- chain ------------------------------------------------------------------

def bs(path, params=None):
    """One Blockscout call. Raises rather than returning partial data."""
    url = BS + path + ("?" + urllib.parse.urlencode(params) if params else "")
    delay = 2
    for _ in range(6):
        try:
            req = urllib.request.Request(url, headers=HDR)
            data = json.load(urllib.request.urlopen(req, timeout=30))
            time.sleep(0.35)
            return data
        except urllib.error.HTTPError as exc:
            if exc.code == 404:
                return None
            time.sleep(delay)
            delay = min(60, delay * 2)
        except Exception:
            time.sleep(delay)
            delay = min(60, delay * 2)
    raise RuntimeError(f"explorer kept failing on {path}; not writing a partial log")


def fetch_wallet(refresh):
    if not refresh and os.path.exists(WALLET_CACHE):
        return json.load(open(WALLET_CACHE, encoding="utf-8"))
    txs, params = [], None
    while True:
        page = bs(f"/addresses/{WALLET}/transactions", params) or {}
        txs += page.get("items", [])
        params = page.get("next_page_params")
        print(f"  {len(txs)} transactions", end="\r", flush=True)
        if not params:
            break
    print()
    w = WALLET.lower()
    for i, t in enumerate(txs, 1):
        h = t["hash"]
        if t.get("token_transfers") is None or t.get("token_transfers_overflow"):
            tt = bs(f"/transactions/{h}/token-transfers") or {}
            t["token_transfers"] = tt.get("items", [])
        if (t.get("from") or {}).get("hash", "").lower() == w and \
                (t.get("method") or "").lower() != "approve":
            it = bs(f"/transactions/{h}/internal-transactions") or {}
            t["_internal"] = it.get("items", [])
        print(f"  detail {i}/{len(txs)}", end="\r", flush=True)
    print()
    os.makedirs(os.path.dirname(WALLET_CACHE), exist_ok=True)
    json.dump(txs, open(WALLET_CACHE, "w", encoding="utf-8"))
    return txs


def addr(obj):
    # Blockscout calls it "hash" on accounts but "address_hash" on tokens; reading
    # only "hash" made every token "" and silently disabled the QUOTES filter.
    obj = obj or {}
    return (obj.get("hash") or obj.get("address_hash") or "").lower()


# One clock for the whole workbook: this PC's, which the bot's log lines use.
# The chain and the ledgers' entry_time/exit_time are UTC while the ledgers' ts
# is local, so a paper row could open in UTC and close in local time, 5h30 apart.
TZ = "UTC" + time.strftime("%z")[:3] + ":" + time.strftime("%z")[3:]


def local(ts):
    """Any timestamp -> 'YYYY-MM-DD HH:MM:SS' local. Naive ones are already local."""
    if not ts:
        return ""
    try:
        d = datetime.fromisoformat(str(ts).replace("Z", "+00:00"))
    except ValueError:
        return str(ts).replace("T", " ")[:19]
    return (d.astimezone() if d.tzinfo else d).strftime("%Y-%m-%d %H:%M:%S")


def bot_log_index(paths):
    """Which tx hashes the bot logged, and which ticker goes with which contract.
    The explorer knows a transaction reverted; only the log knows what for."""
    logged, ticker = set(), {}
    for p in paths:
        text = dd.ANSI.sub("", open(p, encoding="utf-8", errors="replace").read())
        logged |= {h.lower() for h in HEX_TX.findall(text)}
        for m in CALL_ADDR.finditer(text):
            ticker[m.group(2).lower()] = m.group(1)
    return logged, ticker


def tx_rows(txs, logged, ticker, swap_tok):
    w = WALLET.lower()
    rows = []
    for t in sorted(txs, key=lambda t: (t.get("block_number") or 0, t.get("position") or 0)):
        ours = addr(t.get("from")) == w
        ok = t.get("status") == "ok"
        method = t.get("method") or ""
        h = t["hash"].lower()
        tt = t.get("token_transfers") or []
        got_all = [x for x in tt if addr(x.get("to")) == w]
        gave_all = [x for x in tt if addr(x.get("from")) == w]
        got = [x for x in got_all if addr(x.get("token")) not in QUOTES]
        gave = [x for x in gave_all if addr(x.get("token")) not in QUOTES]
        value = int(t.get("value") or 0) / 1e18
        eth_in = sum(int(i.get("value") or 0) for i in t.get("_internal", [])
                     if addr(i.get("to")) == w and i.get("success", True)) / 1e18
        if not ok:
            action = "FAILED"
        elif method.lower() == "approve":
            action = "APPROVE"
        elif gave:
            action = "SELL"
        elif got:
            action = "BUY"
        # Only quote tokens moved: ETH -> USDG (the first leg of a two-step buy)
        # is a BUY of USDG, and USDG -> ETH a SELL. Checked after the memecoin
        # legs so a memecoin bought WITH USDG is not read as a USDG sell.
        elif got_all:
            action, got = "BUY", got_all
        elif gave_all:
            action, gave = "SELL", gave_all
        elif not ours and value > 0:
            action = "DEPOSIT"
        elif ours and value > 0:
            action = "TRANSFER OUT"
        elif ours and addr(t.get("to")) == w:
            action = "SELF-SEND"
        else:
            action = method or "OTHER"
        leg = (gave or got or [None])[0]
        sym, amt, tok_addr = "", None, addr((leg or {}).get("token"))
        if leg:
            tok = leg.get("token") or {}
            sym = tok.get("symbol") or ""
            total = leg.get("total") or {}
            try:
                amt = int(total.get("value") or 0) / 10 ** int(total.get("decimals") or 18)
            except (TypeError, ValueError):
                amt = None
        fee = int(((t.get("fee") or {}).get("value")) or 0) / 1e18 if ours else 0.0
        rate = t.get("historic_exchange_rate") or t.get("exchange_rate")
        reason = t.get("revert_reason") or ""
        if isinstance(reason, dict):
            # {"raw": null} is Blockscout for "reverted with no reason string"
            reason = reason.get("raw") or (json.dumps(reason)[:200] if len(reason) > 1 else "")
        gas, limit = int(t.get("gas_used") or 0), int(t.get("gas_limit") or 0)
        if not ok:
            # A failed swap moves no tokens, so name the token from the bot's own
            # swap line for this hash, or from its contract inside the calldata.
            raw = (t.get("raw_input") or "").lower()
            tok_addr = swap_tok.get(h) or next((a for a in ticker if a[2:] in raw), "")
            sym = ticker.get(tok_addr, "")
            oog = limit and gas >= 0.97 * limit
            reason = (f"{t.get('result') or 'reverted'}: {reason or 'the contract gave no reason'}. "
                      f"Gas used {gas:,} of {limit:,} limit"
                      + (" -- ran out of gas." if oog else ", so not out of gas.")
                      + (" The ETH attached was returned; only the gas was spent."
                         if value > 0 else ""))
        elif action == "SELF-SEND":
            reason = ("0 ETH to your own wallet: the usual way to cancel a stuck "
                      "pending transaction")
        if not ours:
            record = ""
        elif h in logged:
            record = "logged"
        elif addr(t.get("to")) in NOT_BOT_ROUTES:
            record = f"not the bot ({NOT_BOT_ROUTES[addr(t.get('to'))]} route, not in its code)"
        else:
            record = "not logged"
        rows.append({
            "time": local(t.get("timestamp")),
            "nonce": t.get("nonce") if ours else None,
            "dir": "out" if ours else "in",
            "action": action, "token": sym, "amount": amt,
            "eth_sent": value if ours and ok else 0.0, "eth_recv": eth_in if ours else value,
            "status": "success" if ok else "failed",
            "reason": dd.redact(str(reason or ""))[:300],
            "gas": int(t["gas_used"]) if t.get("gas_used") else None,
            "fee": fee, "rate": float(rate) if rate else None,
            "hash": t["hash"], "method": method,
            "to": (t.get("to") or {}).get("name") or addr(t.get("to")),
            "record": record, "token_addr": tok_addr,
        })
    return rows


def wallet_now():
    """Balances right now, so the file can prove its own ETH flows add up."""
    a = bs(f"/addresses/{WALLET}") or {}
    toks = []
    for t in bs(f"/addresses/{WALLET}/token-balances") or []:
        tok = t.get("token") or {}
        try:
            amt = int(t.get("value") or 0) / 10 ** int(tok.get("decimals") or 18)
        except (TypeError, ValueError):
            amt = None
        rate = tok.get("exchange_rate")
        toks.append({"symbol": tok.get("symbol") or "?", "amount": amt,
                     "address": (tok.get("address_hash") or tok.get("address") or "").lower(),
                     "price": float(rate) if rate else None})
    return {"eth": int(a.get("coin_balance") or 0) / 1e18,
            "rate": float(a.get("exchange_rate") or 0) or None, "tokens": toks}


def settle_with_chain(trades, txs, st):
    """The log cannot see past the bot; the chain can. A position the log still
    calls open may have been sold from a wallet app (HASHSTR was), and a
    "pending" buy may never have reached the chain at all."""
    held = {str(p.get("ticker") or "").upper() for p in st.get("open_positions", [])}
    for r in trades:
        if r["outcome"] not in ("open", "pending"):
            continue
        tk = r["token"].upper()
        later = [x for x in txs if x["action"] in ("BUY", "SELL") and x["time"] >= r["opened"]
                 and (x["token"] or "").upper() == tk]
        sells = [x for x in later if x["action"] == "SELL"]
        if r["outcome"] == "pending" and not any(x["action"] == "BUY" for x in later):
            r["outcome"], r["reason"] = "never filled", "no buy of this token on chain"
        elif r["outcome"] == "open" and tk not in held and sells:
            s = sells[-1]
            via = "from a wallet app" if s["record"].startswith("not the bot") else "by the bot"
            r["outcome"], r["sell_tx"] = "sold outside the bot", s["hash"]
            r["reason"] = f"sold {via} at {s['time'][:16]}; the bot recorded no exit and no P&L"


# --- bot records ------------------------------------------------------------

HEX_TX = re.compile(r"0x[0-9a-fA-F]{64}")


def is_fixture(ev, live):
    """Rows the test suite wrote into a real ledger. By signature, not by date."""
    if str(ev.get("ticker")) == "TEST":
        return True
    # The signal-queue tests use message 1. Real channel message ids here run in
    # the thousands; without this, 10 queue fixtures slipped through carrying no
    # ticker, contract or tx hash to recognise them by.
    if ev.get("message_id") == 1 or str(ev.get("signal_id") or "").endswith(":1"):
        return True
    ca = str(ev.get("contract_address") or "").lower()
    if len(ca) == 42 and len({ca[i:i + 2] for i in range(2, 42, 2)}) == 1:
        return True                                   # 0x1111..., 0xabab...
    if live:
        for k in ("tx_hash", "tx_hash_buy", "tx_hash_sell"):
            h = ev.get(k)
            if h and not HEX_TX.fullmatch(str(h)):
                return True                           # "0xfirst"
    return False


def ledger(path, live):
    evs = dd.read_events(path) if hasattr(dd, "read_events") else None
    if evs is None:
        from trade_ledger import read_events
        evs = read_events(path)
    keep = [e for e in evs if not is_fixture(e, live)]
    return keep, len(evs) - len(keep)


def session_starts(paths):
    starts = []
    rx = re.compile(r"^(\d{4}-\d\d-\d\d) (\d\d:\d\d:\d\d).*DRY_RUN=(True|False)")
    for p in paths:
        with open(p, encoding="utf-8", errors="replace") as fh:
            for raw in fh:
                m = rx.match(dd.ANSI.sub("", raw))
                if m:
                    starts.append((f"{m.group(1)} {m.group(2)}", m.group(3) == "True"))
    return sorted(starts)


def mode_of(rec, starts):
    key = f"{rec['day']} {rec['time'][:8]}"
    dry = None
    for k, d in starts:
        if k <= key:
            dry = d
        else:
            break
    return "paper" if dry else "live"


def trade_row(t, mode, source):
    rungs = ", ".join(f"{x.get('rung')} @ {float(x.get('multiplier') or 0):.2f}x"
                      for x in t.get("tranches") or [])
    outcome = {"buy_failed": "buy failed"}.get(t.get("outcome"), t.get("outcome") or "")
    closed = t.get("outcome") == "closed"
    return {
        "mode": mode, "source": source,
        "opened": f"{t.get('day', '')} {t.get('opened_at', '')}".strip(),
        "closed": f"{t.get('closed_day', '')} {t.get('closed_at', '')}".strip() if closed else "",
        "token": t.get("ticker") or "", "ca": t.get("contract_address") or "",
        "stake": float(t.get("stake_usd") or 0) or None, "outcome": outcome,
        "result": ("win" if t.get("is_win") else "loss") if closed else "",
        "pnl": float(t.get("pnl_usd") or 0) if closed else None,
        "gas": float(t.get("gas_cost_usd") or 0) or None,
        "rungs": rungs, "reason": dd.redact(t.get("failure_reason") or "")[:200],
        "buy_tx": t.get("tx_hash_buy") or "", "sell_tx": t.get("tx_hash_sell") or "",
        "peak": float(t.get("peak_multiplier") or 0) or None,
        "flag": FLAGS.get(t.get("ticker") or "", ""),
    }


def paper_ledger_rows(evs):
    quality = {}
    for e in evs:
        if e.get("kind") == "entry_quality":
            quality[(e.get("ticker"), str(e.get("contract_address")).lower())] = e.get("spot_over_entry")
    thin = {e.get("ticker") for e in evs if e.get("kind") == "thin_pool_exit"}
    unpr = {e.get("ticker") for e in evs if e.get("kind") == "unpriceable_position"}
    closed_ids = set()
    rows = []
    for e in evs:
        if e.get("kind") != "position_closed":
            continue
        closed_ids.add(e.get("id"))
        tk = e.get("ticker") or ""
        tp = [n for n, k in (("TP1", "tp1_hit"), ("TP2", "tp2_hit"), ("TP3", "tp3_hit")) if e.get(k)]
        flag = FLAGS.get(tk, "")
        if tk in thin:
            flag = flag or "Thin pool: exited at once (AD-035)"
        if tk in unpr:
            flag = flag or "Unpriceable: force-closed (AD-034)"
        rows.append({
            "mode": "paper", "source": "paper ledger",
            "opened": local(e.get("entry_time")),
            "closed": local(e.get("exit_time") or e.get("ts")),
            "token": tk, "ca": e.get("contract_address") or "",
            "stake": float(e.get("stake_usd") or 0) or None, "outcome": "closed",
            "result": "win" if e.get("is_win") else "loss",
            "pnl": float(e.get("pnl_usd") or 0),
            "gas": float(e.get("gas_cost_usd") or 0) or None,
            "rungs": ", ".join(tp), "reason": "",
            "buy_tx": "", "sell_tx": "",
            "peak": float(e.get("peak_multiplier") or 0) or None,
            "quality": quality.get((tk, str(e.get("contract_address")).lower())),
            "flag": flag,
        })
    # A paper position opened in the ledger but never closed was wiped by a paper
    # state reset (the ledger is append-only; the state is not). Calling it
    # "open" would leave it looking live forever.
    still_open = {p.get("id") for p in json.load(
        open(os.path.join(HERE, "cache", "strategy_state_paper.json"), encoding="utf-8")
    ).get("open_positions", [])}
    for e in evs:
        if e.get("kind") == "position_opened" and e.get("id") not in closed_ids:
            rows.append({"mode": "paper", "source": "paper ledger",
                         "opened": local(e.get("entry_time")),
                         "closed": "", "token": e.get("ticker") or "",
                         "ca": e.get("contract_address") or "",
                         "stake": float(e.get("stake_usd") or 0) or None,
                         "outcome": "open" if e.get("id") in still_open
                         else "never closed (state reset)",
                         "result": "", "pnl": None, "gas": None,
                         "rungs": "", "reason": "", "buy_tx": "", "sell_tx": "",
                         "peak": None, "quality": None, "flag": ""})
    rows.sort(key=lambda r: r["opened"])
    return rows


def signal_rows(records, mode_fn):
    label = {"signal": "call received", "skip": "rejected",
             "skip_no_ca": "rejected: no contract address", "buy_failed": "buy failed"}
    out = []
    for r in records:
        if r["kind"] not in label:
            continue
        reason = r.get("reason") or ("no contract address in the call"
                                     if r["kind"] == "skip_no_ca" else "")
        out.append({"mode": mode_fn(r), "time": f"{r['day']} {r['time'][:8]}",
                    "event": label[r["kind"]], "token": r.get("ticker") or "",
                    "detail": r.get("dex") or "", "reason": dd.redact(reason.strip())[:200],
                    "source": "bot log"})
    return out


def safety_rows(evs, mode):
    label = {"thin_pool_exit": "safety exit: thin pool",
             "unpriceable_position": "safety exit: unpriceable",
             "execution_needs_review": "execution needs review"}
    out = []
    for e in evs:
        if e.get("kind") in label:
            out.append({"mode": mode, "time": local(e.get("ts")),
                        "event": label[e["kind"]], "token": e.get("ticker") or "",
                        "detail": "", "reason": dd.redact(str(e.get("reason") or
                                                              e.get("impact_pct") or ""))[:200],
                        "source": "trade ledger"})
    return out


# --- workbook ---------------------------------------------------------------

def style_sheet(ws, headers, widths, formats=None):
    hdr_fill = PatternFill("solid", fgColor="1F3864")
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        cell.fill = hdr_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for c, wdt in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = wdt
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30
    if ws.max_row > 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            if cell.font is None or cell.font.name != FONT:
                cell.font = Font(name=FONT, size=10,
                                 color=cell.font.color if cell.font and cell.hyperlink else None,
                                 underline="single" if cell.hyperlink else None)
            fmt = (formats or {}).get(cell.column)
            if fmt:
                cell.number_format = fmt


def link(ws, row, col, tx):
    if tx and HEX_TX.fullmatch(tx):
        c = ws.cell(row=row, column=col, value=tx)
        c.hyperlink = EXPLORER + tx
        c.font = Font(name=FONT, size=10, color="0563C1", underline="single")
    elif tx:
        ws.cell(row=row, column=col, value=tx)


def build(txs, live_trades, paper_trades, sigs, meta):
    wb = Workbook()
    wb.calculation.fullCalcOnLoad = True     # Excel/Sheets compute every formula on open

    # --- Live Transactions ---
    ws = wb.active
    ws.title = "Live Transactions"
    H = [f"Time ({TZ})", "Nonce", "Direction", "Action", "Token", "Token amount",
         "ETH sent", "ETH received", "Status", "Failure reason / note", "Gas used",
         "Fee (ETH)", "ETH price ($)", "Fee ($)", "Transaction (explorer)", "Method", "To",
         "Bot record", "Net ETH", "Net ($)", "Token contract"]
    for i, r in enumerate(txs, 2):
        vals = [r["time"], r["nonce"], r["dir"], r["action"], r["token"], r["amount"],
                r["eth_sent"], r["eth_recv"], r["status"], r["reason"], r["gas"],
                r["fee"], r["rate"]]
        for c, v in enumerate(vals, 1):
            ws.cell(row=i, column=c, value=v)
        ws.cell(row=i, column=14, value=f'=IF(M{i}="",0,L{i}*M{i})')
        link(ws, i, 15, r["hash"])
        ws.cell(row=i, column=16, value=r["method"])
        ws.cell(row=i, column=17, value=r["to"])
        ws.cell(row=i, column=18, value=r["record"])
        ws.cell(row=i, column=19, value=f"=H{i}-G{i}-L{i}")
        ws.cell(row=i, column=20, value=f'=IF(M{i}="",0,S{i}*M{i})')
        ws.cell(row=i, column=21, value=r["token_addr"])
    style_sheet(ws, H, [19, 7, 9, 11, 14, 16, 11, 12, 9, 50, 10, 12, 11, 10, 68, 22, 28, 30,
                        12, 10, 44],
                {6: "#,##0.00", 7: ETH_FMT, 8: ETH_FMT, 11: "#,##0", 12: ETH_FMT,
                 13: MONEY, 14: MONEY, 19: ETH_FMT, 20: MONEY})

    # --- Wallet Now ---
    ws = wb.create_sheet("Wallet Now")
    wal = meta["wallet"]
    held = {str(p.get("ticker") or "").upper() for p in meta["open_positions"]}
    bought = {x["token_addr"] for x in txs if x["action"] == "BUY"}

    def status(t):
        if t["symbol"].upper() in held:
            return "open position (the bot still tracks it)"
        if (t["amount"] or 0) < 1e-6:
            return "dust left after selling"
        if t["address"] in bought:
            return "bought, still held"
        return "received, never bought (airdrop)"

    assets = [("ETH", "(native)", wal["eth"], wal["rate"], "gas and trading balance")]
    assets += [(t["symbol"], t["address"], t["amount"], t["price"], status(t))
               for t in sorted(wal["tokens"],
                               key=lambda t: -((t["amount"] or 0) * (t["price"] or 0)))]
    for i, (a, ca, amt, px, note) in enumerate(assets, 2):
        for c, v in enumerate([a, ca, amt, px], 1):
            ws.cell(row=i, column=c, value=v)
        ws.cell(row=i, column=5, value=f'=IF(D{i}="",0,C{i}*D{i})')
        ws.cell(row=i, column=6, value=note)
    last = len(assets) + 1
    tot = last + 2
    ws.cell(row=tot, column=1, value="Total ($)")
    ws.cell(row=tot, column=5, value=f"=SUM(E2:E{last})")
    ws.cell(row=tot, column=6, value="unpriced tokens (no explorer price) count as 0")
    ws.cell(row=tot + 1, column=1, value="Total (ETH)")
    ws.cell(row=tot + 1, column=5, value=f"=IFERROR(E{tot}/D2,0)")
    style_sheet(ws, ["Asset", "Contract", "Amount", "Price now ($)", "Value now ($)", "Status"],
                [14, 44, 22, 14, 14, 44],
                {3: "#,##0.000000", 4: '$#,##0.00######', 5: MONEY})
    ws.cell(row=tot + 1, column=5).number_format = ETH_FMT
    for rr in (tot, tot + 1):
        ws.cell(row=rr, column=1).font = Font(name=FONT, size=10, bold=True)
    ws.auto_filter.ref = None

    # --- trades (shared layout for live and paper) ---
    def trades_sheet(title, rows, paper):
        ws = wb.create_sheet(title)
        H = ["#", "Source", f"Opened ({TZ})", f"Closed ({TZ})", "Token", "Contract", "Outcome", "Result",
             "Stake ($)", "P&L ($)", "Fees ($)", "Take-profits hit", "Peak (x)",
             "Failure reason", "Buy tx", "Sell tx", "Entry fill vs spot", "Flag"]
        for i, r in enumerate(rows, 2):
            vals = [i - 1, r["source"], r["opened"], r["closed"], r["token"], r["ca"],
                    r["outcome"], r["result"], r["stake"], r["pnl"], r["gas"],
                    r["rungs"], r["peak"], r["reason"]]
            for c, v in enumerate(vals, 1):
                ws.cell(row=i, column=c, value=v)
            link(ws, i, 15, r["buy_tx"])
            link(ws, i, 16, r["sell_tx"])
            ws.cell(row=i, column=17, value=r.get("quality"))
            ws.cell(row=i, column=18, value=r["flag"])
        style_sheet(ws, H, [5, 15, 19, 19, 16, 44, 11, 8, 10, 10, 9, 22, 9, 34,
                            20 if paper else 68, 20 if paper else 68, 11, 44],
                    {9: MONEY, 10: MONEY, 11: MONEY, 13: "0.00x", 17: "0.000"})
        return ws

    trades_sheet("Live Trades", live_trades, False)
    trades_sheet("Paper Trades", paper_trades, True)

    # --- Signals & Failures ---
    ws = wb.create_sheet("Signals & Failures")
    H = ["Mode", f"Time ({TZ})", "Event", "Token", "Venue", "Reason", "Source"]
    for i, r in enumerate(sigs, 2):
        for c, v in enumerate([r["mode"], r["time"], r["event"], r["token"],
                               r["detail"], r["reason"], r["source"]], 1):
            ws.cell(row=i, column=c, value=v)
    style_sheet(ws, H, [8, 19, 30, 18, 14, 60, 13])

    # --- Summary (formulas over the sheets above) ---
    ws = wb.create_sheet("Summary", 0)
    T = "'Live Transactions'"
    LT, PT, SF = "'Live Trades'", "'Paper Trades'", "'Signals & Failures'"
    title = ws.cell(row=1, column=1, value="Robinhood Chain bot — trade log")
    title.font = Font(name=FONT, bold=True, size=14)
    ws.cell(row=2, column=1, value=f"Generated {meta['generated']}   ·   wallet {WALLET}")
    ws.cell(row=2, column=1).font = Font(name=FONT, size=9, italic=True)
    # Rows are keyed, and every cross-row formula looks its rows up by key. The
    # first draft hardcoded "=B15/B14" for the win rate while wins were on row 17
    # and closed trades on row 16 -- a clean, error-free sheet with wrong numbers.
    # (key, label, live, paper, note, format); formulas may be callables of `at`.
    spec = [
        ("h1", "ON-CHAIN  (live wallet — every transaction)", None, None, None, "head"),
        ("sent", "Transactions sent", f'=COUNTIF({T}!C:C,"out")', None,
         "count of Direction = out", "int"),
        ("buys", "  buys", f'=COUNTIF({T}!D:D,"BUY")', None, "", "int"),
        ("sells", "  sells", f'=COUNTIF({T}!D:D,"SELL")', None, "", "int"),
        ("appr", "  approvals", f'=COUNTIF({T}!D:D,"APPROVE")', None,
         "token allowance needed before a sell", "int"),
        ("fail", "  FAILED", f'=COUNTIF({T}!D:D,"FAILED")', None,
         "reverted on-chain; the gas is still paid", "int"),
        ("gas_eth", "Gas fees paid (ETH)", f"=SUM({T}!L:L)", None, "sum of Fee (ETH)", "eth"),
        ("gas_usd", "Gas fees paid ($)", f"=SUM({T}!N:N)", None,
         "each fee x the ETH price at the time", "money"),
        ("gas_fail", "  of which on FAILED transactions ($)",
         f'=SUMIF({T}!D:D,"FAILED",{T}!N:N)', None, "", "money"),
        ("dep", "Deposits received (ETH)", f'=SUMIF({T}!D:D,"DEPOSIT",{T}!H:H)', None, "", "eth"),
        ("spent", "ETH spent on buys", f'=SUMIF({T}!C:C,"out",{T}!G:G)', None,
         "failed transactions excluded: a revert returns the ETH", "eth"),
        ("recv", "ETH received from sells", f'=SUMIF({T}!C:C,"out",{T}!H:H)', None, "", "eth"),
        ("implied", "ETH balance this file adds up to",
         lambda at, c: f"={c}{at['dep']}-{c}{at['spent']}+{c}{at['recv']}-{c}{at['gas_eth']}",
         None, "deposits - spent + received - gas", "eth"),
        ("chain", "ETH balance on chain now", meta["wallet"]["eth"], None,
         f"read from the explorer at {meta['generated']}", "eth"),
        ("diff", "  difference", lambda at, c: f"=ROUND({c}{at['chain']}-{c}{at['implied']},12)", None,
         "0 = no ETH movement is missing from this file", "eth"),
        ("notbot", "Sent from a wallet app, not the bot", f'=COUNTIF({T}!R:R,"not the bot*")',
         None, "1inch / 0x routes; the bot's code never calls them", "int"),
        ("unlog", "Sent by the bot but missing from its log",
         f'=COUNTIFS({T}!R:R,"not logged",{T}!D:D,"<>APPROVE")', None,
         "approvals aside, which it never logs by hash", "int"),
        ("b0", None, None, None, None, None),
        ("h_real", "REAL MONEY  (live wallet, deposits to now)", None, None, None, "head"),
        ("now_usd", "Wallet value now ($)", f"='Wallet Now'!E{tot}", None,
         "ETH + priced tokens (see Wallet Now); COG and dust count as 0", "money"),
        ("now_eth", "  in ETH", f"='Wallet Now'!E{tot + 1}", None, "", "eth"),
        ("net_eth", "Net result (ETH)", lambda at, c: f"={c}{at['now_eth']}-{c}{at['dep']}",
         None, "value now - deposits", "eth"),
        ("net_usd", "Net result ($, at today's ETH price)",
         lambda at, c: f"={c}{at['net_eth']}*'Wallet Now'!D2", None,
         "measured in ETH first, so ETH's own price move is not counted", "money"),
        ("bot_pnl", "  the bot's own P&L on live trades ($)", lambda at, c: f"={c}{at['pnl']}",
         None, "only positions it tracked to a close; the gap is gas, leftovers and "
         "trades outside the bot", "money"),
        ("b1", None, None, None, None, None),
        ("h2", "TRADES", "Live", "Paper", None, "head"),
        ("closed", "Positions closed", f'=COUNTIF({LT}!G:G,"closed")',
         f'=COUNTIF({PT}!G:G,"closed")', "", "int"),
        ("wins", "  wins", f'=COUNTIFS({LT}!G:G,"closed",{LT}!H:H,"win")',
         f'=COUNTIFS({PT}!G:G,"closed",{PT}!H:H,"win")', "", "int"),
        ("losses", "  losses", f'=COUNTIFS({LT}!G:G,"closed",{LT}!H:H,"loss")',
         f'=COUNTIFS({PT}!G:G,"closed",{PT}!H:H,"loss")', "", "int"),
        ("wr", "Win rate",
         lambda at, c: f"=IFERROR({c}{at['wins']}/{c}{at['closed']},0)",
         lambda at, c: f"=IFERROR({c}{at['wins']}/{c}{at['closed']},0)", "wins / closed", "pct"),
        ("pnl", "Realised P&L ($)", f'=SUMIF({LT}!G:G,"closed",{LT}!J:J)',
         f'=SUMIF({PT}!G:G,"closed",{PT}!J:J)', "sum of closed-trade P&L", "money"),
        ("avg", "  average per trade ($)",
         lambda at, c: f"=IFERROR({c}{at['pnl']}/{c}{at['closed']},0)",
         lambda at, c: f"=IFERROR({c}{at['pnl']}/{c}{at['closed']},0)", "", "money"),
        # Only "not real" rows. Thin-pool and unpriceable exits are real outcomes
        # of real safety rules; dropping them made paper look $77 better than it was.
        ("clean", "  paper P&L without rows known not to be real ($)", None,
         f'=SUMIFS({PT}!J:J,{PT}!G:G,"closed",{PT}!R:R,"<>*not real*")',
         "drops only rows whose Flag says 'not real' (the price-feed bug)", "money"),
        ("bfail", "Buy attempts that failed",
         f'=COUNTIFS({SF}!A:A,"live",{SF}!C:C,"buy failed")',
         f'=COUNTIFS({SF}!A:A,"paper",{SF}!C:C,"buy failed")', "from the bot log", "int"),
        ("open", "Positions still open", f'=COUNTIF({LT}!G:G,"open")',
         f'=COUNTIF({PT}!G:G,"open")', "", "int"),
        ("b2", None, None, None, None, None),
        ("h3", "SIGNALS", "Live", "Paper", None, "head"),
        ("calls", "Calls received", f'=COUNTIFS({SF}!A:A,"live",{SF}!C:C,"call received")',
         f'=COUNTIFS({SF}!A:A,"paper",{SF}!C:C,"call received")', "", "int"),
        ("rej", "Rejected by filters", f'=COUNTIFS({SF}!A:A,"live",{SF}!C:C,"rejected*")',
         f'=COUNTIFS({SF}!A:A,"paper",{SF}!C:C,"rejected*")',
         "includes calls with no contract address", "int"),
        ("safe", "Safety exits (thin pool / unpriceable)",
         f'=COUNTIFS({SF}!A:A,"live",{SF}!C:C,"safety exit*")',
         f'=COUNTIFS({SF}!A:A,"paper",{SF}!C:C,"safety exit*")', "", "int"),
        ("b3", None, None, None, None, None),
        ("h4", "RECONCILIATION", None, None, None, "head"),
        ("state", "Live state file says", meta["state_line"], None,
         "Source: cache/strategy_state.json (the bot's own counter)", "text"),
        ("fx", "Test rows excluded from the ledgers",
         meta["fixtures_live"], meta["fixtures_paper"],
         "written by the test suite; not real trades (see Notes)", "int"),
    ]
    at = {key: 4 + i for i, (key, *_rest) in enumerate(spec)}
    fmts = {"eth": ETH_FMT, "money": MONEY, "pct": PCT, "int": "#,##0"}
    for key, label, live_v, paper_v, note, kind in spec:
        r = at[key]
        vals = [label,
                live_v(at, "B") if callable(live_v) else live_v,
                paper_v(at, "C") if callable(paper_v) else paper_v,
                note]
        for col, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=col, value=v)
            head = kind == "head"
            cell.font = Font(name=FONT, size=10, bold=head, color="1F3864" if head else None)
            if col in (2, 3) and kind in fmts:
                cell.number_format = fmts[kind]
    for col, wdt in zip("ABCD", (44, 18, 18, 58)):
        ws.column_dimensions[col].width = wdt
    ws.freeze_panes = "A4"

    # --- Notes ---
    ws = wb.create_sheet("Notes")
    notes = [
        ("How to read this file", ""),
        ("Live Transactions", "Every transaction on the live wallet, read from the chain via "
         "the Blockscout explorer. This is the ground truth for fees and failures."),
        ("Bot record", "Whether the bot's own log mentions the transaction's hash. 'not the "
         "bot' = sent through 1inch or 0x, routers nothing in the bot's code calls, so from a "
         "wallet app. 'not logged' covers approvals (never logged by hash), transactions from "
         "before the logs begin, and the first leg of two-step buys (ETH to USDG or a stock "
         "token), whose hash the bot does not log."),
        ("Wallet Now", "Balances read from the explorer when this file was built. GME, GOOGL "
         "and SNOW are stock tokens left over from two-step buys: the bot swapped ETH into the "
         "stock, the second swap (stock into the memecoin) reverted, and a bug in the retry "
         "path ('logger' is not defined -- since fixed) ended the attempt without swapping "
         "back. They still have value and can be sold. COG is the one position the bot still "
         "tracks; the rest is dust or airdrops."),
        ("Live Trades", "One row per position from the bot's log and trade ledger, live "
         "sessions only. 'buy failed' rows are attempts that never filled."),
        ("Paper Trades", "Simulated. No transaction was sent; fills use a real quote plus the "
         "configured slippage, and fees are simulated. Includes every paper session since "
         "09-19 at different stakes and settings, plus DRY_RUN sessions of 09-12."),
        ("Signals & Failures", "Every call received, every rejection with its reason, every "
         "failed buy, and every safety exit, tagged live or paper."),
        ("", ""),
        ("What was excluded", ""),
        ("Test rows", f"{meta['fixtures_live']} rows in the live ledger and "
         f"{meta['fixtures_paper']} in the paper ledger were written by this project's own "
         "test suite (ticker TEST, tx '0xfirst', fixture contracts like 0x1111...). "
         "They are not trades and are not in this file."),
        ("", ""),
        ("Known caveats", ""),
        ("Paper vs live", "Paper results do not include failed transactions, real gas, or "
         "execution delay. Live results are the only ones that cost real money."),
        ("Flagged paper rows", "A Flag marks a row worth a second look. Only a flag saying 'not "
         "real' means the number is wrong -- ONCHAIN_842B9C, a +$448.96 'win' produced by a "
         "price-feed bug -- and only that row is left out of the Summary's second paper P&L "
         "line. Thin-pool and unpriceable exits are real outcomes of safety rules and stay in."),
        ("Live counts", f"The bot's state file reports: {meta['state_line']}. Live Trades has "
         "two more closed trades -- PRPC and AGAINPAD, the first two closes on 09-12, made "
         "before that counter began -- so 1 more win, 1 more loss, and P&L that agrees to the "
         "cent-rounding of the log. Every closed live trade matches a buy and a sell on chain."),
        ("Formulas", "Summary and Fee ($) are formulas; they recalculate when opened in Excel "
         "or Google Sheets."),
    ]
    for i, (a, b) in enumerate(notes, 1):
        ca, cb = ws.cell(row=i, column=1, value=a), ws.cell(row=i, column=2, value=b)
        bold = bool(a) and not b
        ca.font = Font(name=FONT, size=10, bold=True if a else False,
                       color="1F3864" if bold else None)
        cb.font = Font(name=FONT, size=10)
        cb.alignment = Alignment(wrap_text=True, vertical="top")
        ca.alignment = Alignment(vertical="top")
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 110

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    return OUT


# --- safety: nothing secret leaves in a file built to be shared ---------------

def secret_check(path):
    """Fail the build if any .env value appears anywhere in the workbook XML.
    Values are compared, never printed."""
    # Only variables NAMED like secrets. .env also holds public contract
    # addresses, which legitimately appear in the "To" column; comparing every
    # value would abort the build on a false alarm.
    named = re.compile(r"KEY|SECRET|HASH|PASS|RPC|URL|SESSION|API|PHONE", re.I)
    secrets = []
    env = os.path.join(HERE, ".env")
    if os.path.exists(env):
        for line in open(env, encoding="utf-8", errors="replace"):
            if "=" not in line or line.lstrip().startswith("#"):
                continue
            name, v = line.split("=", 1)
            v = v.strip().strip('"').strip("'")
            if not named.search(name) or re.fullmatch(r"0x[0-9a-fA-F]{40}", v):
                continue                           # public address, not a secret
            if v.isdigit() and len(v) < 12:
                continue                           # short ids collide with real numbers
            if len(v) >= 8:
                secrets.append(v)
            # an RPC key can travel without its URL, e.g. inside a copied log line
            secrets += [seg for seg in re.split(r"[/?=&]", v) if len(seg) >= 24]
    with zipfile.ZipFile(path) as z:
        blob = "".join(z.read(n).decode("utf-8", "replace") for n in z.namelist())
    leaked = sum(1 for s in secrets if s in blob)
    host = "quiknode.pro" in blob
    return leaked, host, len(secrets)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--refresh", action="store_true")
    args = ap.parse_args()

    print("bot records...")
    live_evs, fx_live = ledger(LIVE_LEDGER, live=True)
    paper_evs, fx_paper = ledger(PAPER_LEDGER, live=False)

    live_dir_logs = dd.log_files()
    starts = session_starts(live_dir_logs)
    recs = dd.parse_logs(live_dir_logs)
    logged, ticker = bot_log_index(live_dir_logs)
    swap_tok = {r["tx"].lower(): r["address"].lower() for r in recs
                if r["kind"] == "swap" and r.get("tx") and r.get("address")}

    print("live wallet history (on-chain)...")
    txs = tx_rows(fetch_wallet(args.refresh), logged, ticker, swap_tok)
    live_recs = [r for r in recs if mode_of(r, starts) == "live"]
    early_paper_recs = [r for r in recs if mode_of(r, starts) == "paper"]
    paper_recs = dd.parse_logs([PAPER_LOG]) if os.path.exists(PAPER_LOG) else []

    live_trades = [trade_row(t, "live", t.get("source") or "bot log")
                   for t in dd.build_trades(live_recs, live_evs)]
    early = [trade_row(t, "paper", "paper, 09-12 dry run")
             for t in dd.build_trades(early_paper_recs, [])]
    for r in early:          # those sessions ended on 09-12: nothing there is still open
        if r["outcome"] == "open":
            r["outcome"] = "never closed (session ended)"
    paper_trades = sorted(early + paper_ledger_rows(paper_evs), key=lambda r: r["opened"])

    sigs = (signal_rows(recs, lambda r: mode_of(r, starts))
            + signal_rows(paper_recs, lambda r: "paper")
            + safety_rows(live_evs, "live") + safety_rows(paper_evs, "paper"))
    sigs.sort(key=lambda r: r["time"])

    st = json.load(open(os.path.join(HERE, "cache", "strategy_state.json"), encoding="utf-8"))
    state_line = (f"{st.get('total_trades')} trades, {st.get('wins')}W / {st.get('losses')}L, "
                  f"P&L ${float(st.get('total_pnl_usd') or 0):+.2f}, open: "
                  f"{', '.join(p.get('ticker', '?') for p in st.get('open_positions', [])) or 'none'}")
    settle_with_chain(live_trades, txs, st)
    print("wallet balances now...")
    meta = {"generated": datetime.now().strftime("%Y-%m-%d %H:%M ") + TZ,
            "fixtures_live": fx_live, "fixtures_paper": fx_paper, "state_line": state_line,
            "wallet": wallet_now(), "open_positions": st.get("open_positions", [])}

    path = build(txs, live_trades, paper_trades, sigs, meta)

    leaked, host, n = secret_check(path)
    if leaked or host:
        os.remove(path)
        print(f"ABORTED: {leaked} secret value(s) / keyed RPC host found in the workbook. "
              "File deleted rather than shipped.")
        return 1

    c = Counter(r["action"] for r in txs)
    print(f"\nwrote {path}")
    print(f"  live transactions: {len(txs)}  {dict(c.most_common())}")
    print(f"  live trades: {len(live_trades)}  paper trades: {len(paper_trades)}  "
          f"signal/failure rows: {len(sigs)}")
    print(f"  test rows excluded: live {fx_live}, paper {fx_paper}")
    print(f"  secret check: {n} .env values compared, 0 found; no keyed RPC host")
    return 0


if __name__ == "__main__":
    sys.exit(main())
